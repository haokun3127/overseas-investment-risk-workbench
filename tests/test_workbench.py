import asyncio
import copy
import io
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from app.config import Settings
from app.main import create_app
from app.ingest import parse_file, insert_document
from app.model import validate_result, validate_endpoint
from app.retrieval import retrieve
from app.store import Store, DEFAULT_RULES

HEADERS={'X-Requested-With':'risk-workbench'}

def test_preflight_cancellation_and_worker_claim(client):
    doc_id=add(client)
    s=client.app.state.settings;s.model_mode='external';s.model_base_url='';s.model_name='';s.model_api_key=''
    check=client.post('/api/jobs/preflight',json={'document_ids':[doc_id]}).json()
    assert not check['ready'] and len(check['blockers'])==2
    assert check['items'][0]['issues']==['尚未授权外部分析']
    s.model_mode='local';s.model_base_url='http://127.0.0.1:1/v1';s.model_name='test'
    rules=copy.deepcopy(DEFAULT_RULES);rules['confirmed']=True;client.put('/api/rules',json=rules)
    assert client.post('/api/jobs/preflight',json={'document_ids':[doc_id]}).json()['ready']
    jid=client.post('/api/jobs',json={'document_ids':[doc_id]}).json()['items'][0]['id']
    store=client.app.state.store
    stale=store.one('SELECT * FROM jobs WHERE id=?',(jid,))
    assert client.post(f'/api/jobs/{jid}/cancel').status_code==200
    asyncio.run(client.app.state.worker.process(stale))
    assert store.one('SELECT status FROM jobs WHERE id=?',(jid,))['status']=='cancelled'
    assert not store.rows('SELECT * FROM analyses')
    assert client.post(f'/api/jobs/{jid}/cancel').status_code==409
    assert client.post('/api/jobs/999999/cancel').status_code==404
    assert client.get('/api/jobs?status=cancelled').json()['total']==1
    assert client.get('/api/jobs?status=queued').json()['total']==0
    assert client.get('/api/jobs?q='+str(jid)).json()['total']==1
    new=client.post('/api/jobs',json={'document_ids':[doc_id]}).json()['items'][0]['id']
    assert new!=jid

def test_report_scope_and_date_validation(client):
    client.post('/api/demo')
    report=client.get('/api/report?demo=demo&level=red').json()
    assert report['total']==1
    assert report['items'][0]['evidence'] and report['items'][0]['rule_version']
    assert client.get('/api/report?demo=real').json()['total']==0
    for path in ['/api/events','/api/export/events','/api/report']:
        assert client.get(path+'?date_from=2026-09-20&date_to=2026-09-01').status_code==400
        assert client.get(path+'?demo=invalid').status_code==400
        assert client.get(path+'?date_from=invalid').status_code==400

def test_job_pagination_searches_beyond_recent_page(client):
    doc_id=add(client);store=client.app.state.store
    with store.db() as db:
        for _ in range(105):db.execute("INSERT INTO jobs(document_id,status,created_at) VALUES (?,'completed','test')",(doc_id,))
    first=client.get('/api/jobs').json();second=client.get('/api/jobs?page=2').json()
    assert first['total']==105 and len(first['items'])==100 and len(second['items'])==5
    assert client.get('/api/jobs?q=1').json()['total']==1

def test_workspace_filters_export_and_review_history(client):
    local_id=add(client,source='本地业务资料')
    external_id=add(client,title='允许外发的资料',body='这是允许外发的另一份完整测试正文。',allow_external=True)
    workspace=client.get('/api/workspace').json()
    assert workspace['unanalyzed']==2
    assert workspace['documents']['event']['authorized']==1
    assert [r['id'] for r in client.get('/api/documents?policy=external').json()['items']]==[external_id]
    assert [r['id'] for r in client.get('/api/documents?q=本地业务').json()['items']]==[local_id]
    assert client.get('/api/documents?status=unanalysed').json()['total']==2
    client.post('/api/demo')
    event=client.get('/api/events?demo=demo&level=red').json()['items'][0]
    for status,note in [('needs_change','需要核实停电范围'),('accepted','已核对原始资料')]:
        assert client.put(f"/api/events/{event['id']}/review",json={'status':status,'note':note}).status_code==200
    history=client.get(f"/api/events/{event['id']}").json()['review_history']
    assert [r['status'] for r in history]==['accepted','needs_change']
    assert history[1]['note']=='需要核实停电范围'
    response=client.get('/api/export/events?demo=demo&level=red')
    import csv
    rows=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(rows)==2 and rows[1][13]=='虚构演示'
    assert rows[1][8]=='已核对原始资料'
    assert len(list(csv.reader(io.StringIO(client.get('/api/export/events?demo=real').content.decode('utf-8-sig')))))==1
    assert client.put('/api/events/999999/review',json={'status':'accepted'}).status_code==404

def test_export_escapes_spreadsheet_formulas(client):
    client.post('/api/demo')
    event=client.get('/api/events').json()['items'][0]
    client.put(f"/api/events/{event['id']}/review",json={'status':'accepted','note':'=HYPERLINK("bad")'})
    data=client.get('/api/export/events').content.decode('utf-8-sig')
    assert "'=HYPERLINK" in data

def test_model_check_uses_no_business_documents(client,monkeypatch):
    from app import model
    settings=client.app.state.settings
    settings.model_base_url='https://example.test/v1';settings.model_name='test';settings.model_api_key='test'
    async def fake_call(settings,payload):
        assert payload['references']==[]
        assert payload['rules']==DEFAULT_RULES
        assert '合成测试' in payload['document']['title']
        return {'decision':'insufficient','reason':'没有业务资料','events':[]},{}
    monkeypatch.setattr(model,'call_model',fake_call)
    assert client.post('/api/model/check').json()['ok'] is True
    runtime=client.get('/api/runtime').json()
    assert runtime['worker_enabled'] is False and runtime['disk_free_bytes']>0

@pytest.fixture
def client(tmp_path):
    app=create_app(Settings(data_dir=tmp_path),run_worker=False)
    with TestClient(app,headers=HEADERS) as c:
        assert c.post('/api/auth/setup',json={'password':'test-password-2026'}).status_code==200
        yield c

def add(c,**kw):
    data={'title':'供电中断测试资料','body':'某制造业园区供电设施故障，生产线已经暂停，恢复时间未确定。','date':'2026-09-01','source':'合成测试资料',**kw}
    r=c.post('/api/documents',json=data)
    assert r.status_code==200,r.text
    return r.json()['id']

def test_authentication_csrf_and_no_secret(client):
    config=client.get('/api/settings').json()
    assert 'model_api_key' not in json.dumps(config)
    assert client.post('/api/demo',headers={'Origin':'https://evil.example'}).status_code==403
    assert client.post('/api/demo',headers={'X-Requested-With':''}).status_code==403
    assert client.post('/api/auth/setup',json={'password':'another-password'}).status_code==409
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/documents').status_code==401
    assert client.post('/api/auth/login',json={'password':'wrong-password'}).status_code==401
    assert client.post('/api/auth/login',json={'password':'test-password-2026'}).status_code==200

def test_setup_is_local_only(tmp_path):
    app=create_app(Settings(data_dir=tmp_path),False)
    with TestClient(app,headers=HEADERS,client=('192.168.5.5',321)) as c:
        assert c.post('/api/auth/setup',json={'password':'test-password-2026'}).status_code==403

def test_import_csv_partial_and_duplicate(client):
    csv='标题,时间,来源,正文,原始链接\n测试标题,2026-09-01,合成资料,这是用于测试的完整风险资料,\n空正文,2026-09-01,测试,,\n'
    r=client.post('/api/import',files={'file':('samples.csv',csv.encode('utf-8-sig'),'text/csv')})
    assert r.status_code==200,r.text
    assert r.json()['created']==1
    assert r.json()['results'][1]['status']=='failed'
    r2=client.post('/api/import',files={'file':('samples.csv',csv.encode(),'text/csv')})
    assert r2.json()['results'][0]['status']=='duplicate'
    doc_id=r.json()['results'][0]['id']
    assert client.get(f'/api/documents/{doc_id}/file').content==csv.encode('utf-8-sig')

def test_file_parsers():
    doc=Document();doc.add_paragraph('这是正文内容');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='表格';table.cell(0,1).text='依据'
    out=io.BytesIO();doc.save(out)
    assert '依据' in parse_file('test.docx',out.getvalue())[0]['record']['body']
    book=Workbook();book.active.append(['标题','正文']);book.active.append(['案例','资料正文']);out=io.BytesIO();book.save(out)
    assert parse_file('test.xlsx',out.getvalue())[0]['record']['title']=='案例'
    assert parse_file('sample.txt','中文测试正文'.encode())[0]['record']['body']=='中文测试正文'

def test_invalid_imports(client):
    assert client.post('/api/documents',json={'title':'标题','body':'正文','url':'javascript:alert(1)'}).status_code==400
    assert client.post('/api/documents',json={'title':'标题','body':'正文','date':'yesterday'}).status_code==400
    assert client.post('/api/import',files={'file':('broken.pdf',b'not pdf')}).status_code==400
    assert client.post('/api/import',files={'file':('macro.exe',b'hello')}).status_code==400
    assert client.post('/api/import',files={'file':('sample.txt',b'x'*(10*1024*1024+1))}).status_code==413

def test_retrieval_respects_external_permission_and_demo(client):
    local=add(client,title='供电应急知识',body='供电设施故障时需核实停电范围并启用应急供电方案。',kind='knowledge')
    allowed=add(client,title='可外发供电知识',body='供电中断影响生产时，应记录供电恢复计划。',kind='knowledge',allow_external=True)
    store=client.app.state.store
    assert local in {r['document_id'] for r in retrieve(store,'供电设施故障')}
    results=retrieve(store,'供电设施故障',external=True)
    assert local not in {r['document_id'] for r in results}
    assert allowed in {r['document_id'] for r in results}

def test_job_preconditions(client):
    doc_id=add(client)
    assert client.post('/api/jobs',json={'document_ids':[doc_id]}).status_code==409
    s=client.app.state.settings;s.model_base_url='https://test.example/v1';s.model_name='test-model';s.model_api_key='secret'
    assert client.post('/api/jobs',json={'document_ids':[doc_id]}).status_code==409
    rules=copy.deepcopy(DEFAULT_RULES);rules['confirmed']=True
    assert client.put('/api/rules',json=rules).status_code==200
    assert client.post('/api/jobs',json={'document_ids':[doc_id]}).status_code==400
    allowed=add(client,title='授权资料',allow_external=True)
    first=client.post('/api/jobs',json={'document_ids':[allowed]}).json()
    second=client.post('/api/jobs',json={'document_ids':[allowed]}).json()
    assert first['items'][0]['id']==second['items'][0]['id']
    assert second['items'][0]['existing']

def valid_result(text='这是用于测试的风险证据原文。'):
    return {'decision':'risk','reason':'资料中存在经营中断信息','events':[{'title':'测试风险','category':'环境与安全','level':'red','summary':'测试摘要','impact':'测试影响推断','recommendation':'建议核实','evidence':[{'reference':'S1','quote':text,'reason':'支持中断判断'}]}]}

def test_citation_and_classification_validation():
    references={'S1':{'text':'这是用于测试的风险证据原文。','kind':'source','document_id':1,'locator':'正文','title':'测试资料'}}
    assert validate_result(valid_result(),references,DEFAULT_RULES)['events'][0]['evidence'][0]['document_id']==1
    bad=valid_result('不存在的证据引用内容')
    with pytest.raises(ValueError,match='原文匹配'):validate_result(bad,references,DEFAULT_RULES)
    bad=valid_result();bad['events'][0]['category']='不存在的分类'
    with pytest.raises(ValueError,match='分类或等级'):validate_result(bad,references,DEFAULT_RULES)
    bad=valid_result();bad['decision']='non_risk'
    with pytest.raises(ValueError,match='不一致'):validate_result(bad,references,DEFAULT_RULES)

def test_local_mode_endpoint_boundary(tmp_path):
    with pytest.raises(ValueError):validate_endpoint(Settings(data_dir=tmp_path,model_mode='local',model_base_url='https://8.8.8.8/v1'))
    with pytest.raises(ValueError):validate_endpoint(Settings(data_dir=tmp_path,model_mode='external',model_base_url='http://example.com/v1'))
    validate_endpoint(Settings(data_dir=tmp_path,model_mode='local',model_base_url='http://127.0.0.1:1234/v1'))

def test_demo_review_and_filtering(client):
    assert client.post('/api/demo').json()['added']==4
    assert client.post('/api/demo').json()['added']==0
    assert client.get('/api/overview?demo=real').json()['total']==0
    assert client.get('/api/overview?demo=demo').json()['total']==4
    event=client.get('/api/events?level=red').json()['items'][0]
    assert event['demo']==1
    assert client.put(f'/api/events/{event["id"]}/review',json={'status':'accepted','note':'测试复核'}).status_code==200
    detail=client.get(f'/api/events/{event["id"]}').json()
    assert detail['review_note']=='测试复核'
    assert detail['evidence'][0]['quote'] in client.get(f'/api/documents/{detail["document_id"]}').json()['body']
    assert client.get('/api/events?date_from=2026-09-07').json()['total']==2

class FakeModel(BaseHTTPRequestHandler):
    response_kind='risk'
    last_payload=None
    def do_POST(self):
        payload=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        FakeModel.last_payload=payload
        request=json.loads(payload['messages'][1]['content'])
        source=next(r for r in request['references'] if r['kind']=='source')
        result=valid_result(source['text'])
        result['events'][0]['evidence'][0]['reference']=source['reference']
        if self.response_kind=='bad':result['events'][0]['evidence'][0]['quote']='这是没有出现在原始资料中的伪造引用内容'
        if self.response_kind in ('non_risk','insufficient'):result={'decision':self.response_kind,'reason':'合成协议测试结果','events':[]}
        response={'choices':[{'message':{'content':json.dumps(result,ensure_ascii=False)}}],'usage':{'prompt_tokens':123,'completion_tokens':45}}
        body=json.dumps(response,ensure_ascii=False).encode();self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(body)
    def log_message(self,*args):pass

def test_worker_real_http_protocol_and_failed_version_retention(client):
    server=ThreadingHTTPServer(('127.0.0.1',0),FakeModel)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        s=client.app.state.settings;s.model_mode='local';s.model_base_url=f'http://127.0.0.1:{server.server_port}/v1';s.model_name='synthetic-test-model'
        rules=copy.deepcopy(DEFAULT_RULES);rules['confirmed']=True;client.put('/api/rules',json=rules)
        doc_id=add(client)
        def run(kind):
            FakeModel.response_kind=kind
            job_id=client.post('/api/jobs',json={'document_ids':[doc_id]}).json()['items'][0]['id']
            job=client.app.state.store.one('SELECT * FROM jobs WHERE id=?',(job_id,))
            asyncio.run(client.app.state.worker.process(job))
            return client.app.state.store.one('SELECT * FROM jobs WHERE id=?',(job_id,))
        assert run('risk')['status']=='completed'
        assert client.get('/api/events?demo=real').json()['total']==1
        assert FakeModel.last_payload['model']=='synthetic-test-model'
        assert run('bad')['status']=='failed'
        assert client.get('/api/events?demo=real').json()['total']==1
        assert len(client.get(f'/api/documents/{doc_id}').json()['analyses'])==1
        assert run('non_risk')['status']=='completed'
        assert client.get('/api/events?demo=real').json()['total']==0
        versions=client.get(f'/api/documents/{doc_id}').json()['analyses']
        assert [a['version'] for a in versions]==[2,1]
        assert versions[0]['decision']=='non_risk'
    finally:
        server.shutdown();server.server_close()

def test_restart_marks_running_job_failed(client):
    doc_id=add(client);store=client.app.state.store
    with store.db() as db:db.execute("INSERT INTO jobs(document_id,status,created_at) VALUES (?,'running','test')",(doc_id,))
    reopened=Store(store.directory)
    assert reopened.one('SELECT status FROM jobs')['status']=='failed'

def test_static_assets_have_local_only_policy(client):
    response=client.get('/')
    assert response.status_code==200
    assert "script-src 'self'" in response.headers['Content-Security-Policy']
    assert client.get('/static/app.js').status_code==200

def test_policy_change_and_running_task_guard(client):
    doc_id=add(client)
    assert client.patch(f'/api/documents/{doc_id}/policy',json={'allow_external':True}).status_code==200
    assert client.get(f'/api/documents/{doc_id}').json()['allow_external']==1
    with client.app.state.store.db() as db:
        db.execute("INSERT INTO jobs(document_id,status,created_at) VALUES (?,'queued','test')",(doc_id,))
    assert client.patch(f'/api/documents/{doc_id}/policy',json={'allow_external':False}).status_code==409

def test_backup_preserves_documents_and_files(client,tmp_path):
    from scripts.manage import backup
    add(client)
    destination=tmp_path.parent/(tmp_path.name+'-backup')
    backup(client.app.state.store.directory,destination)
    restored=Store(destination)
    assert restored.one('SELECT COUNT(*) n FROM documents')['n']==1
    with pytest.raises(ValueError):backup(client.app.state.store.directory,destination)

def test_long_document_failed_part_saves_no_partial_analysis(client,monkeypatch):
    from app import model
    s=client.app.state.settings;s.model_mode='local';s.model_base_url='http://127.0.0.1:1/v1';s.model_name='test'
    rules=copy.deepcopy(DEFAULT_RULES);rules['confirmed']=True;client.put('/api/rules',json=rules)
    doc_id=add(client,body='这是测试用的风险资料正文。'*300)
    calls=[]
    async def fake(settings,payload):
        calls.append(payload)
        if len(calls)>1:raise ValueError('后续片段模拟失败')
        source=payload['references'][0]
        result=valid_result(source['text'][:20]);result['events'][0]['evidence'][0]['reference']=source['reference']
        return result,{}
    monkeypatch.setattr(model,'call_model',fake)
    job_id=client.post('/api/jobs',json={'document_ids':[doc_id]}).json()['items'][0]['id']
    store=client.app.state.store
    asyncio.run(client.app.state.worker.process(store.one('SELECT * FROM jobs WHERE id=?',(job_id,))))
    assert len(calls)==2
    assert store.one('SELECT status FROM jobs WHERE id=?',(job_id,))['status']=='failed'
    assert not client.get(f'/api/documents/{doc_id}').json()['analyses']
