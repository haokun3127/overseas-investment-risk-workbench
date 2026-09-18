import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from .store import now

MAX_TEXT = 60000
ALIASES = {'标题': 'title', '时间': 'date', '日期': 'date', '来源': 'source', '正文': 'body', '正文内容': 'body', '内容': 'body', '原始链接': 'url', '链接': 'url'}

def validate_record(record):
    result = {k: str(record.get(k) or '').strip() for k in ['title', 'date', 'source', 'url', 'body']}
    if not result['title'] or not result['body']:
        raise ValueError('标题和正文不能为空。')
    if len(result['body']) > MAX_TEXT:
        raise ValueError('单条正文超过60000字符，请按章节拆分导入。')
    for key, limit in [('title', 240), ('date', 40), ('source', 300), ('url', 2000)]:
        if len(result[key]) > limit:
            raise ValueError(f'{key} 字段过长。')
    if result['date']:
        from datetime import date
        try:
            date.fromisoformat(result['date'])
        except ValueError:
            raise ValueError('时间须为 YYYY-MM-DD 格式，未知时可留空。')
    if result['url'] and (urlparse(result['url']).scheme not in ('http', 'https') or not urlparse(result['url']).netloc):
        raise ValueError('原始链接仅支持完整的 http 或 https 地址。')
    return result

def chunks_for(text, size=2400):
    # Non-overlapping stable offsets preserve evidence and avoid duplicated source coverage.
    return [(f'正文字符 {i+1}–{min(i+size,len(text))}', text[i:i+size]) for i in range(0, len(text), size)]

def parse_file(filename, data):
    ext = Path(filename).suffix.lower()
    if ext in ('.docx', '.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError('解压后的文档超过50MB限制，请拆分。')
    if ext in ('.csv', '.xlsx'):
        if ext == '.csv':
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = data.decode('gb18030')
            rows = list(csv.reader(io.StringIO(text)))
        else:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                rows = []
                for row in book.active.iter_rows(values_only=True):
                    if len(rows) > 500:
                        raise ValueError('一次最多导入500条记录，请拆分表格。')
                    rows.append(list(row))
            finally:
                book.close()
        if not rows:
            raise ValueError('表格为空。')
        headers = [ALIASES.get(str(h).strip(), str(h).strip()) for h in rows[0]]
        if 'title' not in headers or 'body' not in headers:
            raise ValueError('缺少“标题/title”或“正文/body”列，请下载导入模板。')
        if len(set(headers)) != len(headers):
            raise ValueError('列名重复，请检查表头。')
        records = []
        for index, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            mapped = dict(zip(headers, row))
            if hasattr(mapped.get('date'), 'date'):
                mapped['date'] = mapped['date'].date().isoformat()
            records.append({'record': mapped, 'row': index})
        if len(records) > 500:
            raise ValueError('一次最多导入500条记录，请拆分表格。')
        return records
    if ext == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('加密 PDF 无法读取，请先解除加密。')
        if len(reader.pages) > 200:
            raise ValueError('PDF 超过200页，请拆分。')
        body = '\n\n'.join(f'[第{i+1}页]\n{page.extract_text() or ""}' for i, page in enumerate(reader.pages))
        if len(re.sub(r'\[第\d+页\]|\s', '', body)) < 10:
            raise ValueError('PDF 没有足够的可提取文本；扫描件请先进行 OCR。')
    elif ext == '.docx':
        doc = Document(io.BytesIO(data))
        body = '\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            body += '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for r in table.rows)
    elif ext in ('.txt', '.md'):
        try:
            body = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            body = data.decode('gb18030')
    else:
        raise ValueError('支持 .csv、.xlsx、.docx、文本型 .pdf、.txt 和 .md 文件。')
    return [{'record': {'title': Path(filename).stem, 'body': body}, 'row': 1}]

def insert_document(store, record, kind='event', allow_external=False, demo=False, file_path='', filename=''):
    r = validate_record(record)
    digest = hashlib.sha256((kind + '\0' + str(demo) + '\0' + r['title'] + '\0' + r['body']).encode()).hexdigest()
    with store.db() as db:
        existing = db.execute('SELECT id FROM documents WHERE digest=?', (digest,)).fetchone()
        if existing:
            return existing['id'], False
        doc_id = db.execute('''INSERT INTO documents(title,date,source,url,body,kind,allow_external,demo,digest,file_path,filename,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''', (r['title'],r['date'],r['source'],r['url'],r['body'],kind,int(allow_external),int(demo),digest,file_path,filename,now())).lastrowid
        for locator, content in chunks_for(r['body']):
            db.execute('INSERT INTO chunks(document_id,locator,text) VALUES (?,?,?)', (doc_id,locator,content))
        return doc_id, True
